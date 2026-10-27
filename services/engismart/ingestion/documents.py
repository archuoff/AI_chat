"""Parsers for PDF, Word, PowerPoint, plain text / markdown and CSV."""
from __future__ import annotations

import re
import zipfile
from collections import Counter
from pathlib import Path

from .models import Element, ParsedDocument, RawImage
from .tables import dedupe_merged, table_to_elements
from .utils import clean_text, normalize_value, render_fields

_PAGE_NO = re.compile(r"^(page\s*)?\d+(\s*(of|/)\s*\d+)?$|^page$", re.IGNORECASE)
_DOCX_ARTIFACT = re.compile(r"^Comment\d+$")
_HEADING_LIKE = re.compile(r"^[A-Z0-9][^.:;,!?]{0,68}$")


def _media_count(path: Path, prefix: str) -> int:
    with zipfile.ZipFile(path) as z:
        return sum(1 for n in z.namelist() if n.startswith(prefix))


# ---------------------------------------------------------------- PDF
def parse_pdf(path: Path) -> ParsedDocument:
    """Docling (layout + table models) when available; pypdf otherwise. A fallback is always reported as a warning."""
    from ..config import get_settings

    engine = get_settings().pdf_engine
    if engine == "pypdf":
        return _parse_pdf_pypdf(path)
    try:
        from .pdf_docling import parse_pdf_docling

        return parse_pdf_docling(path)
    except Exception as exc:  # noqa: BLE001 - Docling missing, model folder wrong, conversion failed
        if engine == "docling":
            raise
        doc = _parse_pdf_pypdf(path)
        doc.warnings.append(f"Docling was not used ({type(exc).__name__}: {str(exc)[:160]}); read with pypdf: no tables or headings")
        return doc


def _parse_pdf_pypdf(path: Path) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [clean_text(p.extract_text() or "") for p in reader.pages]
    # drop running headers/footers: lines repeated on most pages
    line_freq = Counter(line for p in pages for line in set(p.split("\n")) if line)
    repeated = {ln for ln, n in line_freq.items() if len(pages) >= 4 and n >= 0.4 * len(pages) and len(ln) <= 120}
    elements, empty, images = [], 0, []
    for i, text in enumerate(pages, start=1):
        lines = [ln for ln in text.split("\n") if ln not in repeated and not _PAGE_NO.match(ln.strip())]
        body = clean_text("\n".join(lines))
        try:
            for img in reader.pages[i - 1].images:
                ext = img.name.rsplit(".", 1)[-1].lower() if "." in img.name else "png"
                images.append(RawImage(img.data, ext, {"page": i}, body[:800]))
        except Exception:  # noqa: BLE001 - unusual image encodings: the page text is still indexed
            pass
        if len(body) < 20:
            empty += 1
            continue
        elements.append(Element("text", body, [], {"page": i}))
    warnings = [f"{empty} page(s) had no extractable text (scanned/images?)"] if empty else []
    return ParsedDocument(path.name, "pdf", elements, {"pages": len(pages), "removed_header_lines": len(repeated)}, warnings, images)


# ---------------------------------------------------------------- Word
def _docx_heading_level(p) -> int | None:
    name = (p.style.name if p.style is not None else "").lower()
    if name.startswith("heading"):
        digits = re.findall(r"\d+", name)
        return int(digits[0]) if digits else 1
    if name == "title":
        return 1
    return None


def parse_docx(path: Path) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    elements: list[Element] = []
    section: list[str] = []
    buf: list[str] = []
    prev_was_heading = False
    table_no = 0
    images: list[RawImage] = []
    last_text = ""

    def flush():
        if buf:
            elements.append(Element("text", "\n".join(buf), list(section), {"section": " > ".join(section)}))
            buf.clear()

    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = clean_text(p.text)
            for data, ext in _docx_pictures(child, d):  # pictures with the section + nearest paragraph as context
                where = " > ".join(section)
                context = "\n".join(x for x in (where, text or last_text) if x)[:800]
                images.append(RawImage(data, ext, {"section": where} if where else {}, context))
            if not text or _DOCX_ARTIFACT.match(text):
                continue
            level = _docx_heading_level(p)
            if level is None and _HEADING_LIKE.match(text) and len(text.split()) <= 8:
                level = 0  # unstyled heading guessed from shape (short, no sentence punctuation)
            if level is not None:
                flush()
                if level > 0:
                    section = section[: level - 1] + [text]
                else:  # guessed heading: nest under a heading that directly precedes it, else replace
                    section = (section[-1:] if prev_was_heading else section[:-1][-1:]) + [text]
                prev_was_heading = True
                continue
            buf.append(text)
            last_text = text
            prev_was_heading = False
        elif tag == "tbl":
            flush()
            table_no += 1
            t = Table(child, d)
            rows = [[c.text for c in dedupe_merged(r.cells)] for r in t.rows]
            elements.extend(table_to_elements(rows, list(section), {"section": " > ".join(section)}, table_no))
            prev_was_heading = False
    flush()
    stats = {"tables": table_no, "images": _media_count(path, "word/media/")}
    return ParsedDocument(path.name, "docx", elements, stats, images=images)


_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def _docx_pictures(paragraph_el, d) -> list[tuple[bytes, str]]:
    out = []
    for blip in paragraph_el.iter(_BLIP):
        rid = blip.get(_EMBED)
        part = d.part.related_parts.get(rid) if rid else None
        if part is not None and hasattr(part, "blob"):
            out.append((part.blob, str(part.partname).rsplit(".", 1)[-1].lower()))
    return out


# ---------------------------------------------------------------- PowerPoint
def _walk_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from _walk_shapes(sh.shapes)
        else:
            yield sh


def parse_pptx(path: Path) -> ParsedDocument:
    from pptx import Presentation

    prs = Presentation(str(path))
    elements: list[Element] = []
    images: list[RawImage] = []
    pictures = 0
    for i, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = clean_text(title_shape.text_frame.text) if title_shape is not None and title_shape.has_text_frame else ""
        section = [title] if title else [f"Slide {i}"]
        texts, table_no, slide_pics = [], 0, []
        for sh in _walk_shapes(slide.shapes):
            if sh.shape_type == 13:
                pictures += 1
                try:  # SVG pictures expose their PNG fallback here
                    slide_pics.append((sh.image.blob, sh.image.ext))
                except Exception:  # noqa: BLE001 - linked / broken pictures have no embedded bytes
                    pass
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                table_no += 1
                rows = [[c.text for c in dedupe_merged(list(r.cells))] for r in sh.table.rows]
                elements.extend(table_to_elements(rows, section, {"slide": i}, table_no))
            elif sh.has_text_frame:
                t = clean_text(sh.text_frame.text)
                if t and not _PAGE_NO.match(t):
                    texts.append(t)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = clean_text(slide.notes_slide.notes_text_frame.text)
            if notes:
                texts.append(f"Speaker notes: {notes}")
        if texts:
            elements.append(Element("text", "\n".join(texts), section, {"slide": i}))
        context = "\n".join([title or f"Slide {i}", *texts])[:800]
        images.extend(RawImage(data, ext, {"slide": i}, context) for data, ext in slide_pics)
    stats = {"slides": len(prs.slides), "pictures": pictures}
    return ParsedDocument(path.name, "pptx", elements, stats, images=images)


# ---------------------------------------------------------------- Text / Markdown / CSV
def parse_text(path: Path) -> ParsedDocument:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    elements, section, buf = [], [], []

    def flush():
        if buf and clean_text("\n".join(buf)):
            elements.append(Element("text", clean_text("\n".join(buf)), list(section), {}))
        buf.clear()

    for line in raw.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)", line) if path.suffix.lower() == ".md" else None
        if m:
            flush()
            level = len(m.group(1))
            section = section[: level - 1] + [m.group(2).strip()]
        else:
            buf.append(line)
    flush()
    return ParsedDocument(path.name, path.suffix.lstrip(".").lower(), elements)


def parse_csv(path: Path) -> ParsedDocument:
    import pandas as pd

    df = pd.read_csv(path)
    elements = []
    for idx, row in df.iterrows():
        fields = {str(k): nv for k, v in row.items() if (nv := normalize_value(v)) is not None}
        if fields:
            elements.append(Element("record", render_fields(fields), [path.stem], {"row": int(idx) + 2}, fields, str(next(iter(fields.values())))[:150]))
    return ParsedDocument(path.name, "csv", elements, {"rows": len(df)})
