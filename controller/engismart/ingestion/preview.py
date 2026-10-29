"""Excel preview for the Manage page: the original sheet grid + what the parser decided for every cell."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from .excel import analyze_sheet
from .incell import read_incell_pictures

MAX_CELL_CHARS = 220
MAX_RECORDS_PER_TABLE = 300


def _display(v: Any) -> str:
    s = str(v)
    return s if len(s) <= MAX_CELL_CHARS else s[: MAX_CELL_CHARS - 1] + "…"


def _pictures(doc_id: str | None, sheet: str, sh, max_rows: int, max_cols: int) -> list[dict]:
    """The document's pictures that sit on this sheet, with the cell (or merged area) they belong to."""
    if not doc_id:
        return []
    from ..images import catalog  # lazy: the catalog imports the ingestion package

    out = []
    for img in catalog.list_images(doc_id):
        loc = img.location or {}
        m = re.match(r"^([A-Z]+)(\d+)$", str(loc.get("cell", "")))
        if loc.get("sheet") != sheet or not m:
            continue
        c, r = column_index_from_string(m.group(1)), int(m.group(2))
        r, c = sh.anchor.get((r, c), (r, c))  # inside a merged area: drawn in its top-left cell (covered cells are not shown)
        if r > max_rows or c > max_cols:
            continue
        r1, c1 = sh.span.get((r, c), (r, c))
        out.append({"id": img.id, "r": r, "c": c, "r1": min(r1, max_rows), "c1": min(c1, max_cols),
                    "placement": loc.get("placement", "floating"), "status": img.status, "reason": img.reason,
                    "description": img.description, "context": img.context, "width": img.width, "height": img.height})
    return out


def build_excel_preview(path: Path, options: dict, sheet: str | None, max_rows: int, max_cols: int,
                        doc_id: str | None = None) -> dict:
    wb = load_workbook(path, data_only=True)
    try:
        visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
        hidden = [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
        if not visible:
            raise KeyError("workbook has no visible sheets")
        ws = next((w for w in visible if w.title == sheet), None) if sheet else visible[0]
        if ws is None:
            raise KeyError(f"sheet '{sheet}' not found")

        sh, elements, report = analyze_sheet(ws, options, read_incell_pictures(path).get(ws.title))

        # per-cell status from the parser's decisions
        status: dict[tuple[int, int], str] = {}
        for t in report["tables"]:
            if t["skipped"]:
                continue
            c0, r0, c1, r1 = range_boundaries(t["range"])
            role = {r: "title" for r in t["preamble_rows"]}
            role.update({r: "header" for r in t["header_rows"]})
            role.update({r: "section" for r in t["section_rows"]})
            role.update({r: "data" for r in t["data_rows"]})
            for (r, c) in sh.orig:
                if r0 <= r <= r1 and c0 <= c <= c1 and r in role:
                    status[(r, c)] = role[r]
        missed = {(m["cell"]) for m in report["coverage"]["missed"]}
        skipped_tables = [t for t in report["tables"] if t["skipped"]]

        rows_out: dict[int, list[dict]] = {}
        for (r, c), raw in sh.raw.items():
            if r > max_rows or c > max_cols:
                continue
            a = f"{get_column_letter(c)}{r}"
            if (r, c) not in sh.orig:
                s = "ignored"
            elif a in missed:
                s = "missed"
            else:
                s = status.get((r, c), "missed")
            if report["skipped"] or any(_inside(r, c, t["range"]) for t in skipped_tables):
                s = "skipped"
            v = sh.orig.get((r, c), raw)
            rows_out.setdefault(r, []).append({"c": c, "v": _display(v), "s": s, "full": len(str(v)) > MAX_CELL_CHARS})

        merges = []
        for (r0, c0), (r1, c1) in sh.span.items():
            if r0 <= max_rows and c0 <= max_cols:
                merges.append({"r0": r0, "c0": c0, "r1": min(r1, max_rows), "c1": min(c1, max_cols)})

        # records per table, exactly as indexed
        records_by_table: dict[str, list[dict]] = {}
        for el in elements:
            if el.kind != "record":
                continue
            row = el.location["row"]
            col = range_boundaries(el.location["range"])[0]  # side-by-side tables share rows: match by column too
            for t in report["tables"]:
                if row in t["data_rows"] and _inside(row, col, t["range"]):
                    bucket = records_by_table.setdefault(t["id"], [])
                    if len(bucket) < MAX_RECORDS_PER_TABLE:
                        bucket.append({"row": row, "section": el.section[1:], "fields": el.fields})
                    break

        return {
            "file_name": path.name,
            "sheets": [{"name": w.title} for w in visible],
            "hidden_sheets": hidden,
            "sheet": {
                "name": ws.title,
                "max_row": sh.max_row, "max_col": sh.max_col,
                "shown_rows": min(sh.max_row, max_rows), "shown_cols": min(sh.max_col, max_cols),
                "truncated": sh.max_row > max_rows or sh.max_col > max_cols,
                "rows": [{"r": r, "cells": sorted(cells, key=lambda x: x["c"])} for r, cells in sorted(rows_out.items())],
                "merges": merges,
                **{k: v for k, v in report.items() if k not in ("max_row", "max_col")},
                "records": records_by_table,
                "record_count": report["records"],
                "pictures": _pictures(doc_id, ws.title, sh, max_rows, max_cols),
            },
        }
    finally:
        wb.close()


def _inside(r: int, c: int, rng: str) -> bool:
    c0, r0, c1, r1 = range_boundaries(rng)
    return r0 <= r <= r1 and c0 <= c <= c1
