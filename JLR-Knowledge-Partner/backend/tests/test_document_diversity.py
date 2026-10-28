"""One document that copies the content of others must not take every place in the evidence."""
from app.rag import retriever


def unit(doc, n, grade=2):
    return {"chunks": [{"id": f"{doc}{n}", "doc_id": doc, "kind": "record", "chunk_index": n}], "score": 1 - n / 100, "sim": 0.5, "grade": grade}


def docs_of(units):
    return [u["chunks"][0]["doc_id"] for u in units]


def test_other_documents_keep_places_and_no_document_takes_more_than_its_share():
    ranked = [unit("xlsx", i) for i in range(12)] + [unit("docx", 0), unit("docx", 1), unit("docx", 2), unit("pptx", 0)]
    out = retriever._diversify(ranked, 10, min_per_doc=2, max_share=0.6)
    d = docs_of(out)
    assert len(out) == 10 and d.count("docx") >= 2 and d.count("pptx") == 1 and d.count("xlsx") <= 6
    assert [u["chunks"][0]["id"] for u in out][:4] == ["xlsx0", "xlsx1", "xlsx2", "xlsx3"]  # rank order is kept


def test_a_weaker_match_never_pushes_out_a_stronger_one():
    ranked = [unit("xlsx", i, grade=2) for i in range(10)] + [unit("docx", 0, grade=1)]
    out = retriever._diversify(ranked, 10, min_per_doc=2, max_share=0.6)
    assert docs_of(out).count("docx") == 0 or docs_of(out).count("xlsx") >= 6  # grade 1 is only a filler, never a guaranteed place
    assert docs_of(retriever._diversify(ranked[:10], 10, 2, 0.6)) == ["xlsx"] * 10  # one document only: unchanged


def test_when_other_documents_are_few_the_free_places_are_still_used():
    ranked = [unit("xlsx", i) for i in range(12)] + [unit("docx", 0)]
    out = retriever._diversify(ranked, 10, min_per_doc=2, max_share=0.6)
    assert len(out) == 10 and docs_of(out).count("docx") == 1 and docs_of(out).count("xlsx") == 9


class FakeStore:
    def __init__(self, first):
        self.first, self.calls = first, []

    def search(self, query, top_k=8, filters=None, mode="hybrid"):
        self.calls.append(filters or {})
        if (filters or {}).get("!doc_id"):
            return [{"id": "d1", "doc_id": "docx", "kind": "text", "text": "guideline", "score": 0.03, "location": {}}]
        return self.first


def _collect(plan, filters, named):
    """All searches of the question, then the second pass (what the chat graph does across its parallel branches and merge step)."""
    merged = retriever.merge_hits(retriever.run_specs(retriever.search_specs(plan, filters, named)))
    for extra in retriever._second_pass(merged, plan, filters):
        merged = retriever.merge_hits([list(merged.values()), extra])
    return merged


def test_a_document_that_fills_the_pool_triggers_a_second_search_without_it(monkeypatch):
    first = [{"id": f"x{i}", "doc_id": "xlsx", "kind": "record", "text": "row", "score": 0.05 - i / 1000, "location": {}} for i in range(30)]
    store = FakeStore(first)
    monkeypatch.setattr(retriever, "get_store", lambda: store)
    cands = _collect({"standalone_question": "bond flange?", "search_queries": ["bond flange"]}, {"domain": ["adas"]}, [])
    assert "d1" in cands and store.calls[-1] == {"domain": ["adas"], "!doc_id": ["xlsx"]}
    mixed = [dict(h, doc_id="xlsx" if i % 2 else "docx") for i, h in enumerate(first)]
    store2 = FakeStore(mixed)
    monkeypatch.setattr(retriever, "get_store", lambda: store2)
    _collect({"standalone_question": "q", "search_queries": ["q"]}, None, [])
    assert not any(c.get("!doc_id") for c in store2.calls)  # two documents sharing the pool 50 / 50: nothing is crowded out


def test_end_to_end_a_word_document_is_not_crowded_out_by_an_excel_copy_of_its_content(api):
    import io

    import docx
    from openpyxl import Workbook

    from app.rag.retriever import retrieve

    wb = Workbook()
    ws = wb.active
    ws.title = "Rear PDC Sensor"
    ws.append(["Guideline", "Rule"])
    for i in range(80):  # the sheet copies guideline text, one row each, and matches the question strongly
        ws.append([f"minimum bond flange carbon fiber bonnet {i}", f"minimum bond flange carbon fiber bonnet: the minimum bond flange for the carbon fiber bonnet is {10 + i} mm"])
    buf = io.BytesIO()
    wb.save(buf)
    d = docx.Document()
    d.add_heading("Carbon fiber bonnet", level=1)
    d.add_paragraph("The minimum bond flange for the carbon fiber bonnet is 18 mm.")
    dbuf = io.BytesIO()
    d.save(dbuf)

    for name, data in (("ADAS Sensor Study Data.xlsx", buf.getvalue()), ("Design Guidelines.docx", dbuf.getvalue())):
        chk = api.post("/api/documents/check", files={"file": (name, data)}, data={"domain": "adas"}).json()
        done = api.post("/api/documents/commit", json={"check_id": chk["check_id"], "domain": "adas", "intent": "new"})
        assert done.status_code == 200, done.text
    plan = {"standalone_question": "minimum bond flange for the carbon fiber bonnet", "search_queries": ["minimum bond flange carbon fiber bonnet"],
            "sub_questions": [], "intent": "fact", "sheets": [], "entities": [], "ids": []}
    names = [e["file_name"] for e in retrieve(plan["standalone_question"], plan, {"domain": ["adas"], "!superseded": True})]
    assert "Design Guidelines.docx" in names, names
    assert len(names) <= 8  # the docx has one relevant paragraph: it gets its place, the free places stay with the best rows
