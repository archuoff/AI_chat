"""Excel parser: finds every table on every visible sheet and turns each row into a record.

Handles:
  * several tables per sheet (separated by blank rows or blank columns)
  * title / info lines above a table (single-value rows)  -> text element, last one = table title
  * multi-row headers built from merged cells              -> "Part Dimensions > Width (mm)"
  * section rows inside a table (single-value rows)        -> section label on following rows
  * vertically merged data cells (FMEA hierarchy)          -> value filled down to every row
  * Place-in-Cell pictures (stored as #VALUE! + rich data)   -> "[picture: alt text]" in the cell, the picture itself is
                                                                extracted with its column / row / box context
  * stacked tables with no blank row between them           -> a bold or filled title bar followed by a bold / filled
                                                                header row starts a new table (not a "section row")
  * nested section rows                                      -> a wide / bold / filled bar = level 1, a plain label under it = level 2
  * a vertically merged label over a header + picture row    -> the picture row is data, not a second header row
  * other Excel errors (#REF! ...)                           -> ignored (listed in the coverage report)

Every decision is recorded (TableInfo + coverage) so the Manage page can show exactly what was
captured, and a user can correct it with parse options:

    {"sheets": {"<sheet>": {"skip": true}},
     "tables": {"<sheet>!<range>": {"header_rows": 2, "skip": false}}}
"""
from __future__ import annotations

from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .incell import InCellPicture, read_incell_pictures
from .models import Element, ParsedDocument, RawImage
from .utils import is_textual, normalize_value, render_fields, unique_keys

MAX_HEADER_ROWS = 4
MAX_LISTED_CELLS = 200
PICTURE = "[picture"  # marker text of a Place-in-Cell picture ("[picture]" / "[picture: alt]" / "[picture linked: url]")
WIDE_BAR = 0.6  # a single value merged over this share of the table width is a title / section bar


def is_picture(v: Any) -> bool:
    return isinstance(v, str) and v.startswith(PICTURE)


def _header_text(v: Any) -> str:
    return " ".join(str(v).split())


def addr(r: int, c: int) -> str:
    return f"{get_column_letter(c)}{r}"


@dataclass
class TableInfo:
    id: str  # "<sheet>!<range>": stable key for parse options
    range: str
    kind: str  # "table" | "text"
    title: str | None = None
    preamble_rows: list[int] = field(default_factory=list)
    header_rows: list[int] = field(default_factory=list)
    data_rows: list[int] = field(default_factory=list)
    section_rows: list[int] = field(default_factory=list)
    columns: list[dict] = field(default_factory=list)  # [{"col": "A", "name": "Vehicle / Model", "filled": 5}]
    records: int = 0
    header_rows_detected: int = 0
    overridden: bool = False
    skipped: bool = False


class _Sheet:
    def __init__(self, ws, pictures: dict[tuple[int, int], InCellPicture] | None = None):
        self.name = ws.title
        self.orig: dict[tuple[int, int], Any] = {}
        self.raw: dict[tuple[int, int], Any] = {}  # every non-empty cell before cleaning
        self.emph: set[tuple[int, int]] = set()  # bold or solid-filled cells: title bars / headers
        self.pictures = pictures or {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None or (isinstance(c.value, str) and not c.value.strip()):
                    continue
                self.raw[(c.row, c.column)] = c.value
                v = normalize_value(c.value)
                if v is not None:
                    self.orig[(c.row, c.column)] = v
                    if c.font.b or c.fill.fill_type == "solid":
                        self.emph.add((c.row, c.column))
        for (r, c), pic in self.pictures.items():  # Place-in-Cell: the cell holds #VALUE! in the file, but it IS a picture
            self.raw.setdefault((r, c), "#VALUE!")
            self.orig[(r, c)] = (f"[picture linked: {pic.url[:80]}]" if pic.url and not pic.data
                                 else f"[picture: {pic.alt[:80]}]" if pic.alt else "[picture]")
        self.anchor: dict[tuple[int, int], tuple[int, int]] = {}
        self.span: dict[tuple[int, int], tuple[int, int]] = {}
        for rng in ws.merged_cells.ranges:
            a = (rng.min_row, rng.min_col)
            self.span[a] = (rng.max_row, rng.max_col)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.anchor[(r, c)] = a
        self.filled = set(self.orig) | {cell for cell, a in self.anchor.items() if a in self.orig}
        self.max_row = ws.max_row
        self.max_col = ws.max_column
        self.images = len(getattr(ws, "_images", []))

    def val(self, r: int, c: int) -> Any:
        return self.orig.get(self.anchor.get((r, c), (r, c)))

    def h_merged(self, r: int, c: int) -> bool:
        a = self.anchor.get((r, c))
        return a is not None and self.span[a][1] > a[1]


def _runs(sorted_ints: list[int]) -> list[tuple[int, int]]:
    runs: list[tuple[int, int]] = []
    for x in sorted_ints:
        if runs and x == runs[-1][1] + 1:
            runs[-1] = (runs[-1][0], x)
        else:
            runs.append((x, x))
    return runs


def _blocks(sh: _Sheet) -> list[tuple[int, int, int, int]]:
    """Rectangular regions separated by fully blank rows, then by fully blank columns.

    Column groups covering (nearly) the same rows are one table with empty spacer / picture
    columns in between; a group covering different rows is a separate side-by-side table.
    """
    out = []
    for r0, r1 in _runs(sorted({r for r, _ in sh.filled})):
        groups: list[list] = []  # [c0, c1, rows]
        for c0, c1 in _runs(sorted({c for r, c in sh.filled if r0 <= r <= r1})):
            rows = {r for r, c in sh.filled if r0 <= r <= r1 and c0 <= c <= c1}
            if groups and len(rows & groups[-1][2]) / len(rows | groups[-1][2]) >= 0.6:
                groups[-1][1], groups[-1][2] = c1, groups[-1][2] | rows
            else:
                groups.append([c0, c1, rows])
        for c0, c1, rows in groups:
            out.append((min(rows), max(rows), c0, c1))
    return out


def _split_stacked(sh: _Sheet, block: tuple[int, int, int, int]) -> list[tuple[int, int, int, int]]:
    """Tables stacked without a blank row: a title bar + a bold / filled header row starts a new table.

    Needs style evidence (bold or fill) so that a plain section row followed by text rows is never mistaken for it."""
    r0, r1, c0, c1 = block
    width = c1 - c0 + 1

    def own(r):
        return [(c, sh.orig[(r, c)]) for c in range(c0, c1 + 1) if (r, c) in sh.orig]

    def emphasised(r, cells):
        return bool(cells) and sum((r, c) in sh.emph for c, _ in cells) >= 0.6 * len(cells)

    def bar(r):
        cells = own(r)
        if len(cells) != 1 or not is_textual(cells[0][1]) or is_picture(cells[0][1]):
            return False
        c = cells[0][0]
        a = sh.anchor.get((r, c))
        span = sh.span[a][1] - c + 1 if a == (r, c) else 1
        return span >= WIDE_BAR * width or (r, c) in sh.emph

    def header_like(r):
        cells = own(r)
        return (len(cells) >= 2 and all(is_textual(v) and not is_picture(v) and len(v) <= 80 for _, v in cells)
                and len({v for _, v in cells}) == len(cells) and emphasised(r, cells))

    cuts = []
    for r in range(r0 + 1, r1 - 1):
        above_has_table = any(len(own(x)) >= 2 for x in range(r0, r))
        below = own(r + 2)
        if above_has_table and bar(r) and header_like(r + 1) and below and not emphasised(r + 2, below):
            cuts.append(r)
    out, start = [], r0
    for r in cuts:
        out.append((start, r - 1, c0, c1))
        start = r
    out.append((start, r1, c0, c1))
    return out


def _parse_block(sh: _Sheet, r0: int, r1: int, c0: int, c1: int, opts: dict,
                 captured: set) -> tuple[list[Element], TableInfo]:
    cols = range(c0, c1 + 1)
    rows = list(range(r0, r1 + 1))
    rng = f"{addr(r0, c0)}:{addr(r1, c1)}"
    info = TableInfo(id=f"{sh.name}!{rng}", range=rng, kind="table")
    if opts.get("skip"):
        info.skipped = info.overridden = True
        return [], info

    def row_vals(r):
        return [(c, sh.val(r, c)) for c in cols if sh.val(r, c) is not None]

    def single_value(r):
        vals = row_vals(r)
        return len({str(v) for _, v in vals}) == 1 and sum((r, c) in sh.orig for c in cols) <= 1

    def capture_row(r):
        captured.update((r, c) for c in cols if (r, c) in sh.orig)

    elements: list[Element] = []
    i, preamble = 0, []
    while i < len(rows) and single_value(rows[i]):
        preamble.append(str(row_vals(rows[i])[0][1]))
        info.preamble_rows.append(rows[i])
        capture_row(rows[i])
        i += 1
    if preamble:
        elements.append(Element("text", "\n".join(preamble), [sh.name], {"sheet": sh.name, "rows": f"{rows[0]}-{rows[i - 1]}"}))
    if i >= len(rows):
        info.kind = "text"
        return elements, info

    # ---- header detection (or user override)
    h0 = rows[i]
    first = row_vals(h0)
    textual_share = sum(is_textual(v) and not is_picture(v) for _, v in first) / max(len(first), 1)
    n_header = 0

    def header_like_row(r):
        """A row under a vertically merged header cell continues the header only when its own cells are header text:
        a row of pictures, numbers or errors is data (the merged cell is just the label of that row)."""
        cells = [v for c, v in row_vals(r) if (r, c) in sh.orig]
        return not cells or all(is_textual(v) and not is_picture(v) for v in cells)

    if textual_share >= 0.6 and i + 1 < len(rows):
        n_header = 1
        for c in cols:  # merged header cells spanning several rows
            a = sh.anchor.get((h0, c))
            if a == (h0, c):
                end = min(sh.span[a][0], r1)
                if end - h0 + 1 > MAX_HEADER_ROWS:
                    continue  # a label merged over many rows groups the rows below it; it is not a header
                reach = h0
                while reach < end and header_like_row(reach + 1):
                    reach += 1
                n_header = max(n_header, reach - h0 + 1)
        while n_header < MAX_HEADER_ROWS and i + n_header < len(rows) - 1:  # sub-header rows under merged parents
            r = rows[i + n_header]
            cells = [(c, v) for c, v in row_vals(r) if (r, c) in sh.orig]
            under_parent = sum(sh.h_merged(r - 1, c) for c, _ in cells)  # most cells sit under a merged parent; a plain column may join them
            if cells and all(is_textual(v) and not is_picture(v) for _, v in cells) and under_parent >= max(1, 0.5 * len(cells)):
                n_header += 1
            else:
                break
    info.header_rows_detected = n_header
    if opts.get("header_rows") is not None:
        n_header = max(0, min(int(opts["header_rows"]), len(rows) - i - 1))
        info.overridden = True
    header_rows = rows[i : i + n_header]
    info.header_rows = header_rows
    for h in header_rows:
        capture_row(h)
    names, fallbacks = [], []
    for c in cols:
        parts: list[str] = []
        for h in header_rows:
            v = sh.val(h, c)
            if v is not None and (not parts or parts[-1] != str(v)):
                parts.append(str(v))
        names.append(" > ".join(_header_text(p) for p in parts))
        fallbacks.append(f"Column {get_column_letter(c)}")
    headers = dict(zip(cols, unique_keys(names, fallbacks)))
    filled_count = {c: 0 for c in cols}

    # ---- data rows
    table_title = preamble[-1] if preamble else None
    info.title = table_title
    level1 = level2 = None
    level1_wide = False  # section rows: a wide / bold / filled bar = level 1, a plain label under it = level 2
    for r in rows[i + n_header :]:
        if not any((r, c) in sh.orig for c in cols):
            continue  # nothing new on this row (only filled-down merged values)
        if single_value(r) and len(cols) >= 3:
            c_s, v_s = row_vals(r)[0]
            a = sh.anchor.get((r, c_s))
            wide = (a == (r, c_s) and sh.span[a][1] - c_s + 1 >= WIDE_BAR * len(cols)) or (r, c_s) in sh.emph
            if wide or level1 is None or not level1_wide:  # plain rows after a plain row are siblings (old behaviour)
                level1, level2, level1_wide = str(v_s), None, wide
            else:
                level2 = str(v_s)
            info.section_rows.append(r)
            capture_row(r)
            continue
        vals = row_vals(r)
        fields = {headers[c]: v for c, v in vals}
        for c, _ in vals:
            filled_count[c] += 1
        section = [sh.name] + [s for s in (table_title, level1, level2) if s]
        loc = {"sheet": sh.name, "row": r, "range": f"{addr(r, c0)}:{addr(r, c1)}"}
        title = next((v for _, v in vals if is_textual(v)), vals[0][1])  # first text value, not a serial number
        elements.append(Element("record", render_fields(fields), section, loc, fields, str(title)[:150]))
        info.data_rows.append(r)
        capture_row(r)
    info.records = len(info.data_rows)
    info.columns = [{"col": get_column_letter(c), "name": headers[c], "filled": filled_count[c]}
                    for c in cols if filled_count[c] or names[c - c0]]
    if not info.records:
        info.kind = "text"
    return elements, info


def _sheet_report(sh: _Sheet, tables: list[TableInfo], captured: set, skipped: bool) -> dict:
    ignored = sorted(set(sh.raw) - set(sh.orig))  # pictures-in-cells, Excel errors, placeholders
    missed = sorted(set(sh.orig) - captured)
    nonempty = len(sh.orig)
    return {
        "skipped": skipped,
        "dims": f"A1:{addr(max(sh.max_row, 1), max(sh.max_col, 1))}",
        "max_row": sh.max_row, "max_col": sh.max_col,
        "merged_ranges": len(sh.span),
        "images": sh.images,
        "in_cell_pictures": sum(1 for p in sh.pictures.values() if p.data),
        "pictures_unresolved": [{"cell": addr(r, c), "reason": p.problem} for (r, c), p in sorted(sh.pictures.items()) if not p.data][:MAX_LISTED_CELLS],
        "tables": [asdict(t) for t in tables],
        "records": sum(t.records for t in tables),
        "coverage": {
            "nonempty_cells": nonempty,
            "captured_cells": nonempty - len(missed),
            "percent": round(100 * (nonempty - len(missed)) / nonempty, 1) if nonempty else 100.0,
            "missed": [{"cell": addr(r, c), "value": str(sh.orig[(r, c)])[:120]} for r, c in missed[:MAX_LISTED_CELLS]],
            "ignored": [{"cell": addr(r, c), "value": str(sh.raw[(r, c)])[:60]} for r, c in ignored[:MAX_LISTED_CELLS]],
            "ignored_count": len(ignored),
        },
    }


def analyze_sheet(ws, options: dict | None = None, pictures: dict | None = None) -> tuple[_Sheet, list[Element], dict]:
    options = options or {}
    sh = _Sheet(ws, pictures)
    captured: set = set()
    elements: list[Element] = []
    tables: list[TableInfo] = []
    skipped = bool((options.get("sheets") or {}).get(ws.title, {}).get("skip"))
    if not skipped:
        for block in (part for b in _blocks(sh) for part in _split_stacked(sh, b)):
            r0, r1, c0, c1 = block
            key = f"{ws.title}!{addr(r0, c0)}:{addr(r1, c1)}"
            els, info = _parse_block(sh, *block, (options.get("tables") or {}).get(key, {}), captured)
            elements.extend(els)
            tables.append(info)
    return sh, elements, _sheet_report(sh, tables, captured, skipped)


def parse_excel(path: str | Path, options: dict | None = None) -> ParsedDocument:
    path = Path(path)
    wb = load_workbook(path, data_only=True)
    in_cell = read_incell_pictures(path)  # Place-in-Cell pictures: openpyxl shows them as #VALUE!
    elements: list[Element] = []
    stats: dict[str, Any] = {"sheets": {}, "skipped_sheets": []}
    warnings: list[str] = []
    images: list[RawImage] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            stats["skipped_sheets"].append(ws.title)
            continue
        sh, els, report = analyze_sheet(ws, options, in_cell.get(ws.title))
        elements.extend(els)
        stats["sheets"][ws.title] = report
        if not report["skipped"]:
            images.extend(_sheet_pictures(ws, els, report, sh))
            for p in report["pictures_unresolved"]:
                warnings.append(f"Sheet '{ws.title}' {p['cell']}: picture not readable ({p['reason']})")
        cov = report["coverage"]
        if not report["skipped"] and cov["percent"] < 100:
            warnings.append(f"Sheet '{ws.title}': {cov['nonempty_cells'] - cov['captured_cells']} cell(s) not captured "
                            f"({cov['percent']}% coverage)")
    wb.close()
    total = sum(s["coverage"]["nonempty_cells"] for s in stats["sheets"].values() if not s["skipped"])
    got = sum(s["coverage"]["captured_cells"] for s in stats["sheets"].values() if not s["skipped"])
    stats["coverage_percent"] = round(100 * got / total, 1) if total else 100.0
    stats["tables"] = sum(sum(t["kind"] == "table" for t in s["tables"]) for s in stats["sheets"].values())
    return ParsedDocument(path.name, path.suffix.lstrip(".").lower(), elements, stats, warnings, images)


def _picture_context(ws_title: str, sh: _Sheet, report: dict, rows: dict, r: int | None, c: int | None, alt: str = "") -> str:
    """Where a picture sits, in the words of the sheet: box title, column heading, row label, the row's record, alt text.
    These labelled lines are sent to the vision model together with the picture."""
    lines = [f"Sheet: {ws_title}"]
    if r is not None and c is not None:
        for t in report["tables"]:
            r0, r1, c0, c1 = _bounds(t["range"])
            if r0 <= r <= r1 and c0 <= c <= c1:
                if t.get("title"):
                    lines.append(f"Box title: {t['title']}")
                column = next((col["name"] for col in t["columns"] if col["col"] == get_column_letter(c)), "")
                if column:
                    lines.append(f"Column: {column}")
                label = sh.val(r, c0) if c != c0 else None  # first column of the box = the row's label
                if label is not None and not is_picture(label):
                    lines.append(f"Row label: {label}")
                break
    if alt:
        lines.append(f"Alt text: {alt}")
    rec = rows.get(r)
    if rec is not None:
        lines.append("Same row:\n" + rec.text)
    return "\n".join(lines)[:1000]


def _sheet_pictures(ws, elements: list[Element], report: dict, sh: _Sheet | None = None) -> list[RawImage]:
    """Pictures of the sheet, each tied to the cell it sits in: floating ones through their anchor cell, Place-in-Cell
    ones through the cell itself."""
    rows = {e.location.get("row"): e for e in elements if e.kind == "record"}
    sh = sh or _Sheet(ws)
    out = []
    for (r, c), pic in sorted(sh.pictures.items()):
        if not pic.data:
            continue
        loc: dict[str, Any] = {"sheet": ws.title, "row": r, "cell": addr(r, c), "placement": "in_cell"}
        out.append(RawImage(pic.data, pic.ext, loc, _picture_context(ws.title, sh, report, rows, r, c, pic.alt)))
    for img in getattr(ws, "_images", []):
        try:
            data = img._data()
        except Exception:  # noqa: BLE001 - broken / linked picture
            continue
        anchor = getattr(getattr(img, "anchor", None), "_from", None)
        r, c = (anchor.row + 1, anchor.col + 1) if anchor is not None else (None, None)
        loc = {"sheet": ws.title, "placement": "floating"}
        if r is not None:
            loc.update(row=r, cell=addr(r, c))
        out.append(RawImage(data, (getattr(img, "format", None) or "png").lower(), loc, _picture_context(ws.title, sh, report, rows, r, c)))
    return out


def _bounds(rng: str) -> tuple[int, int, int, int]:
    from openpyxl.utils.cell import range_boundaries

    c0, r0, c1, r1 = range_boundaries(rng)
    return r0, r1, c0, c1
