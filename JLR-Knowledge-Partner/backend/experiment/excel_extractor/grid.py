"""The openpyxl pass: cells, styles, merges, sizes, visibility, tables, links, comments, validations, names.

Rule: this module only READS. Values are kept exactly as stored (cached values; dates as ISO text); a merged area's value lives on its
master cell and every covered cell resolves to it through `SheetData.merge_at`.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX

from .model import Cell, Comment, Link, Merge, Sheet, Style, Workbook, a1, rng

ERRORS = {"#VALUE!": "value", "#N/A": "na", "#REF!": "ref", "#DIV/0!": "div0", "#NAME?": "name", "#NULL!": "null", "#NUM!": "num",
          "#CONNECT!": "connect", "#SPILL!": "spill", "#CALC!": "calc", "#GETTING_DATA": "getting_data", "#FIELD!": "field", "#BLOCKED!": "blocked"}
BIG_FILE = 30 * 1024 * 1024  # above this, formulas are not read (second full load) unless asked for


@dataclass
class SheetData:
    """The in-memory working set of one sheet (the JSON model keeps the parts worth saving)."""

    model: Sheet
    cells: dict[tuple[int, int], Cell] = field(default_factory=dict)  # cells that hold a value / formula / picture
    frames: dict[tuple[int, int], int] = field(default_factory=dict)  # empty cells that only carry a fill or a border: style id
    merge_at: dict[tuple[int, int], Merge] = field(default_factory=dict)  # every cell of every merged area (master included)

    def style(self, r: int, c: int) -> Style | None:
        cell = self.cells.get((r, c))
        sid = cell.s if cell is not None else self.frames.get((r, c))
        return self.model.styles[sid] if sid is not None else None

    def value_at(self, r: int, c: int) -> Any:
        """Effective value: a covered merged cell reads its master."""
        m = self.merge_at.get((r, c))
        cell = self.cells.get((m.r0, m.c0) if m else (r, c))
        return None if cell is None else cell.v

    def text_at(self, r: int, c: int) -> str:
        v = self.value_at(r, c)
        return "" if v is None else str(v).strip()


# ------------------------------------------------------------------ colours
def _lum(rgb: str) -> float:
    r, g, b = (int(rgb[i : i + 2], 16) for i in (0, 2, 4))
    return (0.2126 * r + 0.7152 * g + 0.0722 * b) / 255


_THEME_LUM = {0: 1.0, 1: 0.0, 2: 0.9, 3: 0.3, 4: 0.5, 5: 0.5, 6: 0.55, 7: 0.7, 8: 0.45, 9: 0.5}  # lt1 dk1 lt2 dk2 accent1..6 (typical Office)


def _color(c) -> tuple[str | None, float | None]:
    """(label, luminance 0..1) of an openpyxl colour; None when unusable."""
    if c is None:
        return None, None
    try:
        if c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) in (6, 8):
            h = c.rgb[-6:]
            return h.upper(), _lum(h)
        if c.type == "indexed" and c.indexed is not None and int(c.indexed) < len(COLOR_INDEX):
            h = COLOR_INDEX[int(c.indexed)][-6:]
            return h.upper(), _lum(h)
        if c.type == "theme" and c.theme is not None:
            base = _THEME_LUM.get(int(c.theme), 0.5)
            tint = float(c.tint or 0)
            lum = base + tint * (1 - base) if tint > 0 else base * (1 + tint)
            return f"theme:{c.theme}{tint:+.2f}" if tint else f"theme:{c.theme}", max(0.0, min(1.0, lum))
    except (AttributeError, TypeError, ValueError):
        pass
    return None, None


def _fill_class(lum: float | None) -> str:
    if lum is None or lum >= 0.97:
        return "none"
    return "light" if lum >= 0.85 else "mid" if lum >= 0.5 else "dark"


def _style(cell) -> Style:
    fill_label, fill_class = None, "none"
    if cell.fill is not None and cell.fill.fill_type:
        fill_label, lum = _color(cell.fill.fgColor)
        fill_class = _fill_class(lum)
        if fill_class == "none":
            fill_label = None
    b = cell.border
    sides = "".join(ch for ch, side in (("T", b.top), ("R", b.right), ("B", b.bottom), ("L", b.left)) if side is not None and side.style)
    f, al = cell.font, cell.alignment
    color = _color(f.color)[0] if f is not None and f.color is not None else None
    return Style(id=-1, fill=fill_label, fill_class=fill_class, bold=bool(f and f.b), italic=bool(f and f.i),
                 size=float(f.sz) if f is not None and f.sz else None, color=color, halign=al.horizontal if al else None,
                 valign=al.vertical if al else None, wrap=bool(al and al.wrap_text), border=sides, numfmt=cell.number_format or "General")


# ------------------------------------------------------------------ display text
_DEC = re.compile(r"0\.(0+)")


def display(v: Any, fmt: str) -> str | None:
    """How a number reads on screen when the number format changes it (percent, fixed decimals, thousands). None = same as str(v)."""
    if isinstance(v, bool) or not isinstance(v, (int, float)) or fmt in ("General", "@"):
        return None
    f = fmt.split(";")[0]
    m = _DEC.search(f)
    dec = len(m.group(1)) if m else 0
    try:
        if "%" in f:
            return f"{v * 100:,.{dec}f}%" if "," in f else f"{v * 100:.{dec}f}%"
        if m or re.search(r"#,##0|0", f):
            if "e" in f.lower():
                return f"{v:.{dec}E}"
            return f"{v:,.{dec}f}" if "," in f else f"{v:.{dec}f}"
    except (ValueError, OverflowError):
        return None
    return None


def _iso(v: Any) -> str:
    return v.isoformat() if isinstance(v, (dt.datetime, dt.date, dt.time)) else str(v)


# ------------------------------------------------------------------ the pass
def read_workbook(path: str | Path, with_formulas: bool | None = None) -> tuple[Workbook, dict[str, SheetData]]:
    path = Path(path)
    if with_formulas is None:
        with_formulas = path.stat().st_size <= BIG_FILE
    wb_values = load_workbook(path, data_only=True)
    wb_formulas = load_workbook(path, data_only=False) if with_formulas else None
    book = Workbook(file=path.name)
    if not with_formulas:
        book.warnings.append("formulas were not read (large file): use --formulas to read them")
    for name, dn in wb_values.defined_names.items():
        book.defined_names.append({"name": name, "refers_to": str(getattr(dn, "attr_text", ""))})
    out: dict[str, SheetData] = {}
    for i, ws in enumerate(wb_values.worksheets):
        fws = wb_formulas[ws.title] if wb_formulas is not None else None
        data = _read_sheet(ws, fws, i, bool(wb_formulas))
        out[ws.title] = data
        book.sheets.append(data.model)
    wb_values.close()
    if wb_formulas is not None:
        wb_formulas.close()
    return book, out


def _read_sheet(ws, fws, index: int, formulas_read: bool) -> SheetData:
    sh = Sheet(name=ws.title, index=index, state=ws.sheet_state or "visible", max_row=ws.max_row or 0, max_col=ws.max_column or 0,
               dims=ws.dimensions or "")
    data = SheetData(sh)
    fmt = ws.sheet_format
    sh.default_col_w = float(fmt.defaultColWidth) if fmt.defaultColWidth else 9.140625  # 64 px, Excel's default
    sh.default_row_h = float(fmt.defaultRowHeight or 15)

    # sizes + visibility
    for dim in ws.column_dimensions.values():
        lo, hi = dim.min or 0, min(dim.max or dim.min or 0, max(sh.max_col, dim.min or 0) + 50)
        for c in range(lo, hi + 1):
            if dim.width and abs(dim.width - sh.default_col_w) > 0.01:
                sh.col_w[c] = float(dim.width)
            if dim.hidden:
                sh.hidden_cols.append(c)
    for r, dim in ws.row_dimensions.items():
        if dim.height:
            sh.row_h[r] = float(dim.height)
        if dim.hidden:
            sh.hidden_rows.append(r)
        if dim.outlineLevel:
            sh.outline_rows[r] = int(dim.outlineLevel)
    sh.freeze = str(ws.freeze_panes or "")
    sh.autofilter = str(ws.auto_filter.ref or "")

    # merges
    for mr in ws.merged_cells.ranges:
        m = Merge(mr.min_row, mr.min_col, mr.max_row, mr.max_col)
        sh.merges.append(m)
        if (m.r1 - m.r0 + 1) * (m.c1 - m.c0 + 1) <= 200_000:
            for r in range(m.r0, m.r1 + 1):
                for c in range(m.c0, m.c1 + 1):
                    data.merge_at[(r, c)] = m
        else:
            sh.warnings.append(f"merged area {m.range} is huge: not expanded cell by cell")

    # cells
    style_ids: dict[tuple, int] = {}
    hidden_rows, hidden_cols = set(sh.hidden_rows), set(sh.hidden_cols)
    fcells = fws._cells if fws is not None else {}

    def style_id(cell) -> int:
        key = tuple(cell._style)
        sid = style_ids.get(key)
        if sid is None:
            st = _style(cell)
            st.id = sid = len(sh.styles)
            sh.styles.append(st)
            style_ids[key] = sid
        return sid

    for (r, c), cell in ws._cells.items():
        if cell.__class__.__name__ == "MergedCell":
            continue
        v = cell.value
        fcell = fcells.get((r, c))
        formula = fcell.value if fcell is not None and isinstance(fcell.value, str) and fcell.value.startswith("=") else None
        if isinstance(v, str):
            v = v.replace("\r\n", "\n").strip("​")
            if not v.strip():
                v = None
        if v is None and formula is None:
            if cell.has_style:  # a fill or border that is part of the layout (title bars, frames)
                st = style_id(cell)
                if sh.styles[st].fill_class != "none" or sh.styles[st].border:
                    data.frames[(r, c)] = st
            continue
        rec = Cell(r=r, c=c, s=style_id(cell), f=formula)
        st = sh.styles[rec.s]
        if cell.data_type == "e" or (isinstance(v, str) and v in ERRORS):
            rec.t, rec.v = "e", str(v)
            rec.flags.append(f"error:{ERRORS.get(str(v), 'other')}")
        elif isinstance(v, bool):
            rec.t, rec.v = "b", v
        elif isinstance(v, (int, float)):
            rec.t, rec.v = "n", v
            rec.d = display(v, st.numfmt)
        elif isinstance(v, (dt.datetime, dt.date, dt.time)):
            rec.t, rec.v = "d", _iso(v)
        elif v is None:  # a formula whose result was never calculated / saved
            rec.t = "f"
            rec.flags.append("formula_no_cache")
        else:
            rec.t, rec.v = "s", str(v)
        if formulas_read is False and cell.data_type == "f":
            rec.f = str(v)
        if (r, c) in data.merge_at:
            m = data.merge_at[(r, c)]
            rec.flags.append("merged_master")
            rec.merge = m.range
        if r in hidden_rows or c in hidden_cols:
            rec.flags.append("hidden")
        if cell.comment is not None:
            rec.flags.append("comment")
            sh.comments.append(Comment(ref=a1(r, c), author=cell.comment.author or "", text=(cell.comment.text or "").strip()))
        data.cells[(r, c)] = rec

    sh.non_empty = sum(1 for x in data.cells.values() if x.v is not None)
    for h in getattr(ws, "_hyperlinks", []):
        sh.links.append(Link(ref=str(h.ref), target=str(h.target or ""), location=str(h.location or ""), tip=str(h.tooltip or "")))
        key = _first_cell(str(h.ref))
        if key in data.cells:
            data.cells[key].flags.append("link")
    for t in ws.tables.values():
        sh.tables.append({"name": t.name, "ref": t.ref, "header_rows": t.headerRowCount if t.headerRowCount is not None else 1,
                          "columns": [col.name for col in t.tableColumns]})
    for dv in ws.data_validations.dataValidation if ws.data_validations else []:
        sh.validations.append({"ref": str(dv.sqref), "type": dv.type, "formula1": dv.formula1, "formula2": dv.formula2})
    for name, dn in ws.defined_names.items():
        sh.defined_names.append({"name": name, "refers_to": str(getattr(dn, "attr_text", ""))})
    if ws.print_area:
        sh.defined_names.append({"name": "Print_Area", "refers_to": str(ws.print_area)})
    return data


def _first_cell(ref: str) -> tuple[int, int]:
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

    col, row = coordinate_from_string(ref.split(":")[0])
    return row, column_index_from_string(col)
