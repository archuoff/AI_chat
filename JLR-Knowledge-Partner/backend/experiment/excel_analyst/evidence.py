"""Read the complete exported model, not the potentially truncated layout.json."""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl.utils.cell import get_column_letter, range_boundaries


def bounds(ref: str) -> tuple[int, int, int, int]:
    if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d*(?::[A-Z]{1,3}[1-9]\d*)?", ref):
        raise ValueError(f"Invalid A1 reference: {ref}")
    c0, r0, c1, r1 = range_boundaries(ref)
    if not (1 <= c0 <= c1 <= 16384 and 1 <= r0 <= r1 <= 1048576):
        raise ValueError(f"Invalid Excel bounds: {ref}")
    return c0, r0, c1, r1


def contains(outer, inner) -> bool:
    return outer[0] <= inner[0] <= inner[2] <= outer[2] and outer[1] <= inner[1] <= inner[3] <= outer[3]


def overlaps(a, b) -> bool:
    return a[0] <= b[2] and b[0] <= a[2] and a[1] <= b[3] and b[1] <= a[3]


def rect(obj: dict) -> str:
    return f"{get_column_letter(obj['c0'])}{obj['r0']}:{get_column_letter(obj['c1'])}{obj['r1']}"


def address(cell: dict) -> str:
    return f"{get_column_letter(cell['c'])}{cell['r']}"


def encoded(value) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)


@dataclass
class WorkItem:
    id: str
    sheet: str
    scope: str
    cells: dict[str, dict]
    pictures: list[dict]
    payload: dict


@dataclass
class Evidence:
    digest: str
    items: list[WorkItem]
    warnings: list[str]
    skipped_sheets: list[str]


def load_evidence(folder: Path, *, sheet: str | None = None, max_cells: int = 200,
                  max_chars: int = 48000, include_hidden: bool = False) -> Evidence:
    if max_cells < 4 or max_chars < 2000:
        raise ValueError("max_cells must be >= 4 and max_chars >= 2000")
    model_bytes = (folder / "model.json").read_bytes()
    model = json.loads(model_bytes)
    digest = hashlib.sha256(model_bytes)
    items, warnings, skipped = [], list(model.get("warnings", [])), []
    sheets = model.get("sheets", [])
    if sheet is not None and sheet not in {s['name'] for s in sheets}:
        raise ValueError(f"Sheet not found: {sheet}")
    for i, sh in enumerate(sheets):
        name = sh['name']
        safe = ''.join(ch if ch.isalnum() or ch in '-_' else '_' for ch in name)[:40]
        data_bytes = (folder / 'cells' / f'{i:02d}_{safe}.jsonl').read_bytes()
        digest.update(data_bytes)
        if (sheet is not None and name != sheet) or (sh.get('state', 'visible') != 'visible' and not include_hidden):
            skipped.append(name)
            continue
        cells = {}
        for line in data_bytes.decode('utf-8').splitlines():
            if not line.strip():
                continue
            cell = json.loads(line)
            ref = address(cell)
            bounds(ref)
            if ref in cells:
                raise ValueError(f"Duplicate cell in {name}: {ref}")
            cells[ref] = cell
        cells = dict(sorted(cells.items(), key=lambda pair: (pair[1]['r'], pair[1]['c'])))
        local_warnings = list(sh.get('warnings', [])) + sh.get('inventory', {}).get('problems', [])
        warnings.extend(f'{name}: {w}' for w in local_warnings)
        candidates = [g for g in sh.get('regions', []) if not g.get('children')]
        # Unassigned content gets an explicit fallback scope rather than disappearing.
        assigned = set()
        for g in candidates:
            box = bounds(rect(g))
            assigned.update(ref for ref in cells if contains(box, bounds(ref)))
        unassigned = [c for ref, c in cells.items() if ref not in assigned]
        boxes = [bounds(rect(g)) for g in candidates]
        unassigned_objects = [obj['anchor'] for obj in sh.get('pictures', []) + sh.get('shapes', [])
                              if not any(overlaps(box, bounds(rect(obj['anchor']))) for box in boxes)]
        fallback = [{'r0': c['r'], 'r1': c['r'], 'c0': c['c'], 'c1': c['c']} for c in unassigned] + unassigned_objects
        if fallback:
            candidates.append({'id': 'unassigned', 'kind': 'unknown', 'r0': min(c['r0'] for c in fallback),
                               'r1': max(c['r1'] for c in fallback), 'c0': min(c['c0'] for c in fallback),
                               'c1': max(c['c1'] for c in fallback), 'needs_llm': True})
            warnings.append(f'{name}: content outside candidate regions; fallback scope may overlap other regions')
        for j, candidate in enumerate(candidates):
            scope = rect(candidate)
            box = bounds(scope)
            own = {ref: c for ref, c in cells.items() if contains(box, bounds(ref))}
            pics = [p for p in sh.get('pictures', []) if overlaps(box, bounds(rect(p['anchor'])))]
            header_rows = set(candidate.get('header_rows', [])) | set(candidate.get('title_rows', []))
            ordered = sorted(own, key=lambda ref: (own[ref]['r'] not in header_rows, own[ref]['r'], own[ref]['c']))
            # Reserve half the budget for evenly distributed samples, including the end.
            n = min(max_cells, len(ordered))
            priority = ordered[:n // 2]
            preferred = set(priority)
            remaining = [ref for ref in own if ref not in preferred]
            count = n - len(priority)
            sample = [remaining[round(k * (len(remaining) - 1) / max(count - 1, 1))] for k in range(count)] if remaining else []
            picked = list(dict.fromkeys(priority + sample))
            candidate_summary = {k: v for k, v in candidate.items() if k != 'data_rows'}
            candidate_summary['data_row_count'] = len(candidate.get('data_rows', []))
            payload = {'sheet': name, 'scope': scope, 'candidate': candidate_summary,
                       'sheet_state': sh.get('state', 'visible'), 'extraction_warnings': local_warnings,
                       'merges': [m for m in sh.get('merges', []) if overlaps(box, bounds(rect(m)))],
                       'styles': sh.get('styles', []),
                       'pictures': pics,
                       'shapes': [s for s in sh.get('shapes', []) if overlaps(box, bounds(rect(s['anchor'])))],
                       'comments': [c for c in sh.get('comments', []) if contains(box, bounds(c['ref']))],
                       'neighbour_regions': [{'id': g['id'], 'range': rect(g), 'kind': g.get('kind', 'unknown')}
                                             for g in candidates[max(0, j - 2):j + 3] if g is not candidate],
                       'source_cell_count': len(own)}
            while True:
                payload['cells'] = [{**own[ref], 'ref': ref} for ref in picked]
                payload['omitted_cells'] = len(own) - len(picked)
                if len(encoded(payload)) <= max_chars:
                    break
                if not picked:
                    raise ValueError(f'{name}!{scope}: metadata exceeds --max-chars; raise the budget or extract a smaller workbook')
                picked = picked[:len(picked) // 2]
            if own and not picked:
                raise ValueError(f'{name}!{scope}: no cell evidence fits; raise --max-chars')
            items.append(WorkItem(f's{i:02d}-r{j:04d}', name, scope, own, pics, payload))
    return Evidence(digest.hexdigest(), items, warnings, skipped)
